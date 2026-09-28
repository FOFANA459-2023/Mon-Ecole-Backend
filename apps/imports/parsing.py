"""Reading spreadsheets that schools already have (Excel or CSV, French or English headers)."""

import csv
import datetime
import io
import re
import unicodedata

from openpyxl import load_workbook

MAX_ROWS = 5000


class ImportFileError(Exception):
    pass


def normalize(text) -> str:
    """'Date de naissance' → 'datedenaissance'; 'Prénom(s)' → 'prenoms'."""
    text = unicodedata.normalize("NFKD", str(text or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", text.lower())


def read_rows(upload) -> tuple[list[str], list[list]]:
    name = (upload.name or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
        except Exception as exc:  # corrupted or not really an Excel file
            raise ImportFileError("The Excel file could not be read.") from exc
        sheet = workbook.worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        raw = upload.read()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        raise ImportFileError("Upload an Excel (.xlsx) or CSV file.")

    rows = [r for r in rows if any(cell not in (None, "") for cell in r)]
    if not rows:
        raise ImportFileError("The file is empty.")
    if len(rows) - 1 > MAX_ROWS:
        raise ImportFileError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    headers = [str(h or "").strip() for h in rows[0]]
    return headers, rows[1:]


def map_columns(headers: list[str], aliases: dict[str, list[str]]) -> dict[str, int]:
    """Find which column holds each known field, accepting French and English names."""
    lookup = {normalize(alias): key for key, names in aliases.items() for alias in [key, *names]}
    mapping: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = lookup.get(normalize(header))
        if key and key not in mapping:
            mapping[key] = index
    return mapping


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


DATE_FORMATS = ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y"]


def parse_date(value) -> datetime.date | None:
    """Accepts Excel dates, 2012-05-31, 31/05/2012, 31-05-2012. Raises ValueError otherwise."""
    if value in (None, ""):
        return None
    if isinstance(value, datetime.datetime):
        return value.date()
    if isinstance(value, datetime.date):
        return value
    text = cell_text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(text)


def digits(phone: str) -> str:
    return re.sub(r"\D", "", phone or "")
