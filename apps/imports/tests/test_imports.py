import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from apps.enrollments.models import Enrollment
from apps.people.models import Guardian, StaffMember, Student


def xlsx(rows) -> SimpleUploadedFile:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        "eleves.xlsx",
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


HEADERS = [
    "Nom *",
    "Prénom *",
    "Sexe",
    "Date de naissance",
    "Classe",
    "Lien du parent",
    "Nom du parent",
    "Prénom du parent",
    "Téléphone du parent",
]


@pytest.mark.django_db
class TestStudentImport:
    def test_template_is_an_excel_file_with_french_headers(self, school, make_member, client_for):
        response = client_for(make_member(school, "admin_staff"), school).get(
            "/api/v1/imports/students/template/"
        )
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content)).active
        assert sheet["B1"].value == "Nom *"

    def test_check_reports_row_errors_without_saving(self, school, make_class, make_member, client_for):
        make_class("7ème A")
        upload = xlsx(
            [
                HEADERS,
                ["Diallo", "Awa", "F", "31/05/2014", "7ème A", "Mère", "Bah", "Mariama", "620000001"],
                ["", "Sans nom", "X", "pas une date", "9ème Z", "", "", "", ""],
            ]
        )
        response = client_for(make_member(school, "admin_staff"), school).post(
            "/api/v1/imports/students/", {"file": upload}, format="multipart"
        )
        assert response.status_code == 200
        assert response.data["total"] == 2
        assert response.data["valid"] == 1
        assert {e["row"] for e in response.data["errors"]} == {3}
        assert {e["column"] for e in response.data["errors"]} == {
            "Nom",
            "Sexe",
            "Date de naissance",
            "Classe",
        }
        assert not Student.objects.exists()

    def test_commit_creates_students_guardians_and_enrolments(
        self, school, make_class, make_member, client_for
    ):
        class_group = make_class("7ème A")
        upload = xlsx(
            [
                HEADERS,
                ["Diallo", "Awa", "F", "31/05/2014", "7ème A", "Mère", "Bah", "Mariama", "620 00 00 01"],
                ["Diallo", "Ismaël", "M", "2016-01-10", "7ème A", "Mère", "Bah", "Mariama", "620000001"],
                ["Keita", "Fanta", "F", "", "", "", "", "", ""],
            ]
        )
        response = client_for(make_member(school, "admin_staff"), school).post(
            "/api/v1/imports/students/", {"file": upload, "commit": "true"}, format="multipart"
        )
        assert response.status_code == 200, response.data
        assert response.data["created"] == 3
        assert Student.objects.count() == 3
        assert Guardian.objects.count() == 1  # siblings share their mother
        assert Enrollment.objects.filter(class_group=class_group, status="active").count() == 2
        numbers = sorted(Student.objects.values_list("student_number", flat=True))
        assert numbers == ["STU-2026-00001", "STU-2026-00002", "STU-2026-00003"]

    def test_commit_with_errors_saves_nothing(self, school, make_class, make_member, client_for):
        make_class("7ème A", capacity=1)
        upload = xlsx(
            [
                HEADERS,
                ["Diallo", "Awa", "F", "", "7ème A", "", "", "", ""],
                ["Keita", "Fanta", "F", "", "7ème A", "", "", "", ""],
            ]
        )
        response = client_for(make_member(school, "admin_staff"), school).post(
            "/api/v1/imports/students/", {"file": upload, "commit": "true"}, format="multipart"
        )
        assert response.status_code == 400
        assert "complète" in response.data["errors"][0]["message"]
        assert not Student.objects.exists()

    def test_csv_with_semicolons_and_english_headers(self, school, year, make_member, client_for):
        content = b"Last name;First name;Gender\nKollie;James;M\n"
        upload = SimpleUploadedFile("students.csv", content, content_type="text/csv")
        response = client_for(make_member(school, "admin_staff"), school).post(
            "/api/v1/imports/students/", {"file": upload, "commit": "true"}, format="multipart"
        )
        assert response.status_code == 200, response.data
        assert Student.objects.get().last_name == "Kollie"

    def test_missing_required_columns_are_reported(self, school, make_member, client_for):
        upload = xlsx([["Prénom"], ["Awa"]])
        response = client_for(make_member(school, "admin_staff"), school).post(
            "/api/v1/imports/students/", {"file": upload}, format="multipart"
        )
        assert response.data["columns_missing"] == ["Nom"]

    def test_teacher_cannot_import(self, school, make_member, client_for):
        upload = xlsx([HEADERS])
        response = client_for(make_member(school, "teacher"), school).post(
            "/api/v1/imports/students/", {"file": upload}, format="multipart"
        )
        assert response.status_code == 403


@pytest.mark.django_db
def test_staff_import(school, year, make_member, client_for):
    upload = xlsx(
        [
            ["Nom", "Prénom", "Catégorie", "E-mail", "Fonction"],
            ["Camara", "Ibrahima", "Enseignant", "i.camara@test.local", "Professeur de maths"],
            ["Sow", "Aminata", "Administratif", "", "Secrétaire"],
        ]
    )
    response = client_for(make_member(school, "director"), school).post(
        "/api/v1/imports/staff/", {"file": upload, "commit": "true"}, format="multipart"
    )
    assert response.status_code == 200, response.data
    assert sorted(StaffMember.objects.values_list("staff_type", flat=True)) == ["administrative", "teacher"]
