"""Bulk import of students (with guardians and class) and staff from a school's existing lists.

Two passes: `run(commit=False)` checks every row and reports problems without saving anything;
`run(commit=True)` saves everything in one transaction, or nothing if any row is invalid.
"""

import io
from dataclasses import dataclass, field

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.academics.models import ClassGroup
from apps.academics.services import current_year
from apps.audit import services as audit
from apps.core.sequences import format_number, reserve_values
from apps.enrollments.models import Enrollment
from apps.people.models import Guardian, StaffMember, Student, StudentGuardian
from apps.people.services import _numbering_year

from .parsing import cell_text, digits, map_columns, normalize, parse_date, read_rows

MESSAGES = {
    "fr": {
        "required": "Obligatoire.",
        "gender": "Valeur attendue : M ou F.",
        "date": "Date invalide (ex. 31/05/2012).",
        "number_taken": "Ce matricule est déjà utilisé.",
        "number_duplicate": "Matricule en double dans le fichier.",
        "class_unknown": "Classe introuvable pour l'année en cours.",
        "no_year": "Aucune année scolaire en cours : créez-la avant d'affecter des classes.",
        "class_full": "La classe {name} est complète ({cap} places).",
        "guardian_contact": "Indiquez le nom et le téléphone (ou l'e-mail) du parent.",
        "duplicate_student": "Cet élève existe déjà (matricule {number}).",
        "staff_type": "Valeur attendue : enseignant, administratif ou soutien.",
        "email": "Adresse e-mail invalide.",
        "no_enrol_permission": "Vous n'êtes pas autorisé à inscrire des élèves : videz la colonne Classe.",
    },
    "en": {
        "required": "Required.",
        "gender": "Expected M or F.",
        "date": "Invalid date (e.g. 2012-05-31).",
        "number_taken": "This number is already used.",
        "number_duplicate": "Duplicate number in the file.",
        "class_unknown": "Class not found in the current school year.",
        "no_year": "No current school year: create it before assigning classes.",
        "class_full": "Class {name} is full ({cap} places).",
        "guardian_contact": "Give the guardian's name and phone (or email).",
        "duplicate_student": "This student already exists (number {number}).",
        "staff_type": "Expected teacher, administrative or support.",
        "email": "Invalid email address.",
        "no_enrol_permission": "You are not allowed to enrol students: empty the Class column.",
    },
}

GENDERS = {
    "m": "M",
    "masculin": "M",
    "male": "M",
    "homme": "M",
    "garcon": "M",
    "g": "M",
    "boy": "M",
    "f": "F",
    "feminin": "F",
    "female": "F",
    "femme": "F",
    "fille": "F",
    "girl": "F",
}
RELATIONSHIPS = {
    "pere": "father",
    "father": "father",
    "papa": "father",
    "mere": "mother",
    "mother": "mother",
    "maman": "mother",
    "tuteur": "guardian",
    "tutrice": "guardian",
    "guardian": "guardian",
    "autre": "other",
    "other": "other",
}
STAFF_TYPES = {
    "enseignant": "teacher",
    "enseignante": "teacher",
    "professeur": "teacher",
    "teacher": "teacher",
    "administratif": "administrative",
    "administrative": "administrative",
    "admin": "administrative",
    "soutien": "support",
    "support": "support",
    "personneldesoutien": "support",
}


@dataclass
class Column:
    key: str
    fr: str
    en: str
    required: bool = False
    aliases: tuple[str, ...] = ()
    example: str = ""


@dataclass
class ImportResult:
    total: int = 0
    valid: int = 0
    created: int = 0
    errors: list[dict] = field(default_factory=list)
    preview: list[dict] = field(default_factory=list)
    columns_found: list[str] = field(default_factory=list)
    columns_missing: list[str] = field(default_factory=list)

    def as_dict(self):
        return {**self.__dict__, "errors": self.errors[:500]}


class BaseImporter:
    kind = ""
    columns: list[Column] = []

    def __init__(self, school, request=None):
        self.school = school
        self.request = request
        self.language = school.default_language if school.default_language in MESSAGES else "fr"
        self.messages = MESSAGES[self.language]

    def label(self, column: Column) -> str:
        return column.fr if self.language == "fr" else column.en

    def template(self) -> bytes:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = self.sheet_title()
        sheet.append([self.label(c) + (" *" if c.required else "") for c in self.columns])
        sheet.append([c.example for c in self.columns])
        fill = PatternFill("solid", fgColor="1F3A5F")
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
        for index, column in enumerate(self.columns, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = max(14, len(self.label(column)) + 4)
        sheet.freeze_panes = "A2"
        notes = workbook.create_sheet("Instructions" if self.language == "en" else "Mode d'emploi")
        for line in self.instructions():
            notes.append([line])
        notes.column_dimensions["A"].width = 110
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def sheet_title(self) -> str:
        raise NotImplementedError

    def instructions(self) -> list[str]:
        raise NotImplementedError

    def run(self, upload, *, commit: bool) -> ImportResult:
        headers, rows = read_rows(upload)
        mapping = map_columns(headers, {c.key: [c.fr, c.en, *c.aliases] for c in self.columns})
        result = ImportResult(total=len(rows))
        result.columns_found = [self.label(c) for c in self.columns if c.key in mapping]
        result.columns_missing = [self.label(c) for c in self.columns if c.required and c.key not in mapping]
        if result.columns_missing:
            return result

        items = []
        for offset, row in enumerate(rows):
            raw = {key: row[index] if index < len(row) else None for key, index in mapping.items()}
            item, errors = self.validate_row(raw)
            row_number = offset + 2  # spreadsheet row (header is row 1)
            for column_key, message in errors:
                column = next(c for c in self.columns if c.key == column_key)
                result.errors.append({"row": row_number, "column": self.label(column), "message": message})
            if not errors:
                item["_row"] = row_number
                items.append(item)
        self.validate_batch(items, result)
        result.valid = result.total - len({e["row"] for e in result.errors})
        result.preview = [self.preview(item) for item in items[:15]]

        if commit and not result.errors:
            with transaction.atomic():
                result.created = self.save(items)
                audit.record(
                    "import",
                    request=self.request,
                    school=self.school,
                    module=self.kind,
                    summary=f"Imported {result.created} {self.kind} from {getattr(upload, 'name', 'file')}",
                    new={"rows": result.total, "created": result.created},
                )
        return result

    # hooks
    def validate_row(self, raw: dict) -> tuple[dict, list[tuple[str, str]]]:
        raise NotImplementedError

    def validate_batch(self, items: list[dict], result: ImportResult) -> None:
        pass

    def preview(self, item: dict) -> dict:
        return {
            k: (v.isoformat() if hasattr(v, "isoformat") else v)
            for k, v in item.items()
            if not k.startswith("_") and not hasattr(v, "pk")
        }

    def save(self, items: list[dict]) -> int:
        raise NotImplementedError

    # helpers
    def _text(self, raw, key, max_length=200) -> str:
        return cell_text(raw.get(key))[:max_length]

    def _date(self, raw, key, errors):
        try:
            return parse_date(raw.get(key))
        except ValueError:
            errors.append((key, self.messages["date"]))
            return None


class StudentImporter(BaseImporter):
    kind = "students"
    columns = [
        Column(
            "student_number", "Matricule", "Student number", aliases=("numero", "number", "id"), example=""
        ),
        Column("last_name", "Nom", "Last name", True, ("nomdefamille", "surname", "familyname"), "Diallo"),
        Column("first_name", "Prénom", "First name", True, ("prenoms", "givenname", "firstnames"), "Awa"),
        Column("gender", "Sexe", "Gender", aliases=("genre", "sex"), example="F"),
        Column(
            "date_of_birth",
            "Date de naissance",
            "Date of birth",
            aliases=("naissance", "dob", "birthdate"),
            example="31/05/2014",
        ),
        Column("place_of_birth", "Lieu de naissance", "Place of birth", example="Conakry"),
        Column("nationality", "Nationalité", "Nationality", example="Guinéenne"),
        Column("phone", "Téléphone", "Phone", aliases=("tel", "telephone", "phonenumber")),
        Column("address", "Adresse", "Address", example="Ratoma, Conakry"),
        Column("class_name", "Classe", "Class", aliases=("class", "classe"), example="7ème A"),
        Column("enrollment_date", "Date d'inscription", "Enrolment date", aliases=("enrollmentdate",)),
        Column("previous_school", "Établissement précédent", "Previous school", aliases=("ecoleprecedente",)),
        Column(
            "guardian_relationship",
            "Lien du parent",
            "Guardian relationship",
            aliases=("lien",),
            example="Mère",
        ),
        Column("guardian_last_name", "Nom du parent", "Guardian last name", example="Bah"),
        Column("guardian_first_name", "Prénom du parent", "Guardian first name", example="Mariama"),
        Column("guardian_phone", "Téléphone du parent", "Guardian phone", example="+224 620 00 00 00"),
        Column("guardian_email", "E-mail du parent", "Guardian email"),
        Column("guardian_occupation", "Profession du parent", "Guardian occupation", example="Commerçante"),
    ]

    def __init__(self, school, request=None):
        super().__init__(school, request)
        codes = getattr(request, "permission_codes", None)
        self.can_enroll = codes is None or "enrollments.create" in codes
        self.year = current_year(school)
        self.classes = (
            {normalize(c.name): c for c in ClassGroup.objects.filter(school=school, academic_year=self.year)}
            if self.year
            else {}
        )
        self.taken_numbers = {
            normalize(n)
            for n in Student.objects.filter(school=school).values_list("student_number", flat=True)
        }
        self.existing = {
            (normalize(s.first_name), normalize(s.last_name), s.date_of_birth): s.student_number
            for s in Student.objects.filter(school=school).only(
                "first_name", "last_name", "date_of_birth", "student_number"
            )
        }

    def sheet_title(self):
        return "Élèves" if self.language == "fr" else "Students"

    def instructions(self):
        if self.language == "fr":
            return [
                "Une ligne par élève. Les colonnes marquées * sont obligatoires.",
                "Matricule : laissez vide pour qu'il soit attribué automatiquement.",
                "Sexe : M ou F. Dates : JJ/MM/AAAA (ex. 31/05/2014).",
                "Classe : nom exact d'une classe de l'année en cours (ex. 7ème A). "
                "Laissez vide pour ne pas inscrire.",
                "Lien du parent : Père, Mère, Tuteur ou Autre. Un parent déjà connu "
                "(même nom et téléphone) est réutilisé pour les frères et sœurs.",
                "Vérifiez d'abord le fichier : aucune donnée n'est enregistrée tant que "
                "toutes les lignes ne sont pas valides.",
            ]
        return [
            "One row per student. Columns marked * are required.",
            "Student number: leave empty to assign one automatically.",
            "Gender: M or F. Dates: YYYY-MM-DD or DD/MM/YYYY.",
            "Class: exact name of a class in the current school year (e.g. Grade 7A). "
            "Leave empty to skip enrolment.",
            "Guardian relationship: Father, Mother, Guardian or Other. A guardian already known "
            "(same name and phone) is reused for siblings.",
            "Check the file first: nothing is saved until every row is valid.",
        ]

    def validate_row(self, raw):
        errors: list[tuple[str, str]] = []
        item = {
            "student_number": self._text(raw, "student_number", 30),
            "last_name": self._text(raw, "last_name", 100),
            "first_name": self._text(raw, "first_name", 100),
            "place_of_birth": self._text(raw, "place_of_birth", 100),
            "nationality": self._text(raw, "nationality", 60),
            "phone": self._text(raw, "phone", 30),
            "address": self._text(raw, "address", 500),
            "previous_school": self._text(raw, "previous_school", 200),
        }
        for key in ("last_name", "first_name"):
            if not item[key]:
                errors.append((key, self.messages["required"]))
        gender = normalize(raw.get("gender"))
        item["gender"] = GENDERS.get(gender, "")
        if gender and not item["gender"]:
            errors.append(("gender", self.messages["gender"]))
        item["date_of_birth"] = self._date(raw, "date_of_birth", errors)
        item["enrollment_date"] = self._date(raw, "enrollment_date", errors)

        class_name = self._text(raw, "class_name", 50)
        item["class_group"] = None
        if class_name:
            if not self.can_enroll:
                errors.append(("class_name", self.messages["no_enrol_permission"]))
            elif self.year is None:
                errors.append(("class_name", self.messages["no_year"]))
            else:
                item["class_group"] = self.classes.get(normalize(class_name))
                if item["class_group"] is None:
                    errors.append(("class_name", self.messages["class_unknown"]))

        guardian = {
            "relationship": RELATIONSHIPS.get(normalize(raw.get("guardian_relationship")), "guardian"),
            "last_name": self._text(raw, "guardian_last_name", 100),
            "first_name": self._text(raw, "guardian_first_name", 100),
            "phone": self._text(raw, "guardian_phone", 30),
            "email": self._text(raw, "guardian_email", 254),
            "occupation": self._text(raw, "guardian_occupation", 100),
        }
        has_guardian = any(guardian[k] for k in ("last_name", "first_name", "phone", "email"))
        if has_guardian and (not guardian["last_name"] or not (guardian["phone"] or guardian["email"])):
            errors.append(("guardian_last_name", self.messages["guardian_contact"]))
        if guardian["email"] and "@" not in guardian["email"]:
            errors.append(("guardian_email", self.messages["email"]))
        item["guardian"] = guardian if has_guardian else None

        number = normalize(item["student_number"])
        if number and number in self.taken_numbers:
            errors.append(("student_number", self.messages["number_taken"]))
        key = (normalize(item["first_name"]), normalize(item["last_name"]), item["date_of_birth"])
        if item["date_of_birth"] and key in self.existing:
            errors.append(("last_name", self.messages["duplicate_student"].format(number=self.existing[key])))
        return item, errors

    def validate_batch(self, items, result):
        seen: dict[str, int] = {}
        for item in items:
            number = normalize(item["student_number"])
            if number:
                if number in seen:
                    result.errors.append(
                        {
                            "row": item["_row"],
                            "column": self.label(self.columns[0]),
                            "message": self.messages["number_duplicate"],
                        }
                    )
                seen[number] = item["_row"]
        # Class capacity: existing students + those in this file.
        per_class: dict[int, list[dict]] = {}
        for item in items:
            if item["class_group"]:
                per_class.setdefault(item["class_group"].pk, []).append(item)
        for class_id, class_items in per_class.items():
            class_group = class_items[0]["class_group"]
            if class_group.capacity is None:
                continue
            enrolled = Enrollment.objects.filter(class_group_id=class_id, status="active").count()
            for item in class_items[max(0, class_group.capacity - enrolled) :]:
                result.errors.append(
                    {
                        "row": item["_row"],
                        "column": self.label(self.columns[9]),
                        "message": self.messages["class_full"].format(
                            name=class_group.name, cap=class_group.capacity
                        ),
                    }
                )
        rows_with_errors = {e["row"] for e in result.errors}
        items[:] = [i for i in items if i["_row"] not in rows_with_errors]

    def preview(self, item):
        return {
            "row": item["_row"],
            "student_number": item["student_number"],
            "full_name": f"{item['first_name']} {item['last_name']}",
            "gender": item["gender"],
            "date_of_birth": item["date_of_birth"].isoformat() if item["date_of_birth"] else None,
            "class_name": item["class_group"].name if item["class_group"] else "",
            "guardian": (
                f"{item['guardian']['first_name']} {item['guardian']['last_name']}".strip()
                if item["guardian"]
                else ""
            ),
        }

    def save(self, items):
        school, user = self.school, getattr(self.request, "user", None)
        prefix = school.settings.student_number_prefix
        numbering_year = _numbering_year(school)
        missing = [i for i in items if not i["student_number"]]
        if missing:
            first = reserve_values(school, "student", numbering_year, len(missing))
            for offset, item in enumerate(missing):
                item["student_number"] = format_number(prefix, numbering_year, first + offset, width=5)

        students = Student.objects.bulk_create(
            [
                Student(
                    school=school,
                    created_by=user,
                    student_number=i["student_number"],
                    first_name=i["first_name"],
                    last_name=i["last_name"],
                    gender=i["gender"],
                    date_of_birth=i["date_of_birth"],
                    place_of_birth=i["place_of_birth"],
                    nationality=i["nationality"],
                    phone=i["phone"],
                    address=i["address"],
                )
                for i in items
            ]
        )

        # Guardians: reuse by (last name, phone digits) within the school, including siblings in this file.
        known = {
            (normalize(g.last_name), digits(g.phone)): g
            for g in Guardian.objects.filter(school=school).exclude(phone="")
        }
        new_guardians: dict[tuple, Guardian] = {}
        for item in items:
            g = item["guardian"]
            if not g:
                continue
            key = (normalize(g["last_name"]), digits(g["phone"]))
            if g["phone"] and key in known:
                item["_guardian"] = known[key]
            elif g["phone"] and key in new_guardians:
                item["_guardian"] = new_guardians[key]
            else:
                guardian = Guardian(
                    school=school,
                    created_by=user,
                    first_name=g["first_name"],
                    last_name=g["last_name"],
                    phone=g["phone"],
                    email=g["email"],
                    occupation=g["occupation"],
                )
                item["_guardian"] = guardian
                if g["phone"]:
                    new_guardians[key] = guardian
        to_create = [
            i["_guardian"] for i in items if i.get("_guardian") is not None and i["_guardian"].pk is None
        ]
        unique_new = list({id(g): g for g in to_create}.values())
        Guardian.objects.bulk_create(unique_new)

        StudentGuardian.objects.bulk_create(
            [
                StudentGuardian(
                    school=school,
                    created_by=user,
                    student=student,
                    guardian=item["_guardian"],
                    relationship=item["guardian"]["relationship"],
                    is_primary=True,
                    is_financial_contact=True,
                )
                for student, item in zip(students, items, strict=True)
                if item.get("_guardian") is not None
            ]
        )

        today = timezone.localdate()
        Enrollment.objects.bulk_create(
            [
                Enrollment(
                    school=school,
                    created_by=user,
                    student=student,
                    academic_year=item["class_group"].academic_year,
                    class_group=item["class_group"],
                    enrollment_date=item["enrollment_date"] or today,
                    kind=Enrollment.Kind.NEW,
                    previous_school=item["previous_school"],
                )
                for student, item in zip(students, items, strict=True)
                if item["class_group"] is not None
            ]
        )
        return len(students)


class StaffImporter(BaseImporter):
    kind = "staff"
    columns = [
        Column(
            "employee_number", "Matricule", "Staff number", aliases=("numero", "number", "employeenumber")
        ),
        Column("last_name", "Nom", "Last name", True, ("nomdefamille", "surname"), "Camara"),
        Column("first_name", "Prénom", "First name", True, ("prenoms", "givenname"), "Ibrahima"),
        Column("gender", "Sexe", "Gender", aliases=("genre", "sex"), example="M"),
        Column("date_of_birth", "Date de naissance", "Date of birth", aliases=("naissance", "dob")),
        Column("phone", "Téléphone", "Phone", aliases=("tel", "telephone"), example="+224 620 00 00 00"),
        Column("email", "E-mail", "Email", aliases=("courriel", "mail"), example="i.camara@ecole.gn"),
        Column("address", "Adresse", "Address"),
        Column("staff_type", "Catégorie", "Staff type", aliases=("type", "categorie"), example="Enseignant"),
        Column(
            "position",
            "Fonction",
            "Position",
            aliases=("poste", "jobtitle"),
            example="Professeur de mathématiques",
        ),
        Column("qualification", "Diplôme", "Qualification", aliases=("qualification",), example="Licence"),
        Column(
            "specialization", "Spécialité", "Specialization", aliases=("specialite",), example="Mathématiques"
        ),
        Column("employment_date", "Date d'embauche", "Employment date", aliases=("dateembauche", "hiredate")),
    ]

    def __init__(self, school, request=None):
        super().__init__(school, request)
        self.taken_numbers = {
            normalize(n)
            for n in StaffMember.objects.filter(school=school).values_list("employee_number", flat=True)
        }

    def sheet_title(self):
        return "Personnel" if self.language == "fr" else "Staff"

    def instructions(self):
        if self.language == "fr":
            return [
                "Une ligne par membre du personnel. Les colonnes marquées * sont obligatoires.",
                "Matricule : laissez vide pour qu'il soit attribué automatiquement.",
                "Catégorie : Enseignant, Administratif ou Soutien (Enseignant par défaut).",
                "Dates : JJ/MM/AAAA. Aucune donnée n'est enregistrée tant que "
                "toutes les lignes ne sont pas valides.",
            ]
        return [
            "One row per staff member. Columns marked * are required.",
            "Staff number: leave empty to assign one automatically.",
            "Staff type: Teacher, Administrative or Support (Teacher by default).",
            "Dates: YYYY-MM-DD or DD/MM/YYYY. Nothing is saved until every row is valid.",
        ]

    def validate_row(self, raw):
        errors: list[tuple[str, str]] = []
        limits = {
            "employee_number": 30,
            "last_name": 100,
            "first_name": 100,
            "phone": 30,
            "email": 254,
            "address": 500,
            "position": 100,
            "qualification": 150,
            "specialization": 150,
        }
        item = {key: self._text(raw, key, size) for key, size in limits.items()}
        for key in ("last_name", "first_name"):
            if not item[key]:
                errors.append((key, self.messages["required"]))
        gender = normalize(raw.get("gender"))
        item["gender"] = GENDERS.get(gender, "")
        if gender and not item["gender"]:
            errors.append(("gender", self.messages["gender"]))
        staff_type = normalize(raw.get("staff_type"))
        item["staff_type"] = STAFF_TYPES.get(staff_type, "teacher" if not staff_type else "")
        if not item["staff_type"]:
            errors.append(("staff_type", self.messages["staff_type"]))
        if item["email"] and "@" not in item["email"]:
            errors.append(("email", self.messages["email"]))
        item["date_of_birth"] = self._date(raw, "date_of_birth", errors)
        item["employment_date"] = self._date(raw, "employment_date", errors)
        if item["employee_number"] and normalize(item["employee_number"]) in self.taken_numbers:
            errors.append(("employee_number", self.messages["number_taken"]))
        return item, errors

    def validate_batch(self, items, result):
        seen = set()
        for item in items:
            number = normalize(item["employee_number"])
            if number and number in seen:
                result.errors.append(
                    {
                        "row": item["_row"],
                        "column": self.label(self.columns[0]),
                        "message": self.messages["number_duplicate"],
                    }
                )
            seen.add(number)
        rows_with_errors = {e["row"] for e in result.errors}
        items[:] = [i for i in items if i["_row"] not in rows_with_errors]

    def preview(self, item):
        return {
            "row": item["_row"],
            "employee_number": item["employee_number"],
            "full_name": f"{item['first_name']} {item['last_name']}",
            "staff_type": item["staff_type"],
            "position": item["position"],
            "email": item["email"],
        }

    def save(self, items):
        school, user = self.school, getattr(self.request, "user", None)
        prefix = school.settings.employee_number_prefix
        numbering_year = _numbering_year(school)
        missing = [i for i in items if not i["employee_number"]]
        if missing:
            first = reserve_values(school, "employee", numbering_year, len(missing))
            for offset, item in enumerate(missing):
                item["employee_number"] = format_number(prefix, numbering_year, first + offset, width=4)
        StaffMember.objects.bulk_create(
            [
                StaffMember(
                    school=school, created_by=user, **{k: v for k, v in i.items() if not k.startswith("_")}
                )
                for i in items
            ]
        )
        return len(items)


IMPORTERS = {"students": StudentImporter, "staff": StaffImporter}
