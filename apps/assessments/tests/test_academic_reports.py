"""Academic reports (/reports/academic/{key}/): grade report, ranking, subject results."""

import io
from decimal import Decimal as D

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.reports.render import _display

from .conftest import URL, gradebook_of

pytestmark = pytest.mark.django_db
REPORTS = "/api/v1/reports/academic/"


@pytest.fixture
def director(school, make_member, client_for):
    return client_for(make_member(school, "director"), school)


@pytest.fixture
def published(setup, term, maths_marked, director):
    assert director.post(f"{URL}{maths_marked['gradebook'].pk}/publish/").status_code == 200
    return maths_marked


def report(client, key, **params):
    response = client.get(f"{REPORTS}{key}/", params)
    assert response.status_code == 200, response.data
    return response.data


class TestGradeReport:
    def test_the_class_s_published_marks_ranked(self, setup, term, published, director):
        data = report(director, "grades", class_group=setup["class"].pk, term=term.pk)
        marks, subjects = data["sections"]
        maths = setup["maths"].pk
        assert [c["label"] for c in marks["columns"]][3:5] == ["MATH", "Moyenne /20"]
        assert [(r["rank"], r["student"], r[f"s{maths}"]) for r in marks["rows"]] == [
            (1, "Moussa Condé", "16.00"),
            (2, "Awa Diallo", "15.17"),
            (3, "Binta Bah", "11.33"),
        ]
        assert marks["totals"]["average"] == "14.17"
        # French is not published: it is left out and named in the note.
        assert "Pas encore publiées : Français." in marks["note"]
        assert [r["subject"] for r in subjects["rows"]] == ["Mathématiques"]
        assert (subjects["rows"][0]["passed"], subjects["rows"][0]["rate"]) == (3, "100.0")

    def test_the_year_uses_the_annual_rules(self, setup, term, published, director):
        data = report(director, "grades", class_group=setup["class"].pk)
        assert "Année 2026-2027" in data["subtitle"]
        assert [r["average"] for r in data["sections"][0]["rows"]] == ["16.00", "15.17", "11.33"]

    def test_a_class_is_required(self, published, director):
        response = director.get(f"{REPORTS}grades/")
        assert response.status_code == 400 and "class_group" in response.data["fields"]

    def test_marks_print_with_the_school_s_decimals(self, school):
        assert _display(D("15.10"), "mark", "GNF", school, "fr") == "15,10"
        assert _display(D("4.0"), "number", "GNF", school, "en") == "4"


class TestRankingAndSubjects:
    def test_a_level_ranked_together_and_class_by_class(
        self, setup, term, published, director, make_class, make_student, level
    ):
        from apps.enrollments import services as enrollment_services

        other = make_class("7ème B")
        enrollment_services.enroll(make_student("Kadiatou", "Sylla"), other)
        data = report(director, "ranking", term=term.pk, level=level.pk)
        titles = [s["title"] for s in data["sections"]]
        assert titles == [
            "Classes",
            "Classement du niveau 7ème année",
            "Classement — 7ème A",
            "Classement — 7ème B",
        ]
        classes = data["sections"][0]["rows"]
        assert [(r["class"], r["average"], r["rate"]) for r in classes] == [
            ("7ème A", "14.17", "100.0"),
            ("7ème B", None, None),
        ]
        level_rows = data["sections"][1]["rows"]
        assert [(r["rank"], r["class"]) for r in level_rows][:3] == [
            (1, "7ème A"),
            (2, "7ème A"),
            (3, "7ème A"),
        ]
        assert level_rows[-1]["rank"] is None

    def test_subject_results_and_what_is_not_published(self, setup, term, published, director):
        data = report(director, "subjects", term=term.pk)
        results, missing = data["sections"]
        assert [(r["class"], r["subject"], r["average"]) for r in results["rows"]] == [
            ("7ème A", "Mathématiques", "14.17")
        ]
        assert [(r["subject"], r["status"]) for r in missing["rows"]] == [("Français", "En cours de saisie")]

    def test_english_labels_follow_the_user(self, setup, term, published, school, make_member, client_for):
        user = make_member(school, "director", email="en@test.local")
        user.language = "en"
        user.save()
        data = report(client_for(user, school), "ranking", term=term.pk)
        assert data["title"] == "Ranking report" and data["sections"][0]["title"] == "Classes"


class TestAccess:
    def test_the_class_teacher_sees_their_class_only(self, setup, term, published, school, client_for):
        camara = client_for(setup["camara"], school)
        params = {"class_group": setup["class"].pk, "term": term.pk}
        assert camara.get(f"{REPORTS}grades/", params).status_code == 200
        assert camara.get(f"{REPORTS}ranking/", {"term": term.pk}).status_code == 403
        assert camara.get(f"{REPORTS}grades/", {**params, "export": "pdf"}).status_code == 403
        # Mr Barry teaches maths in the class but does not lead it.
        barry = client_for(setup["barry"], school)
        assert barry.get(f"{REPORTS}grades/", params).status_code == 404

    def test_exports_are_files_and_audited(self, setup, term, published, director):
        params = {"class_group": setup["class"].pk, "term": term.pk}
        pdf = director.get(f"{REPORTS}grades/", {**params, "export": "pdf"})
        assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
        xlsx = director.get(f"{REPORTS}grades/", {**params, "export": "xlsx"})
        sheet = load_workbook(io.BytesIO(xlsx.content)).worksheets[0]
        values = [cell.value for cell in sheet[5]]
        assert values[1] == "Moussa Condé" and values[3] == D("16.00")
        assert sheet.cell(row=5, column=4).number_format == "0.00"
        csv = director.get(f"{REPORTS}ranking/", {"term": term.pk, "export": "csv"})
        assert csv.status_code == 200 and "Moussa Condé" in csv.content.decode("utf-8-sig")
        assert AuditLog.objects.filter(module="grades", action="export", entity_type="report").count() == 3

    def test_unknown_report(self, director, published):
        assert director.get(f"{REPORTS}nothing/").status_code == 404


def test_a_term_of_another_year_is_refused(setup, term, school, published, director):
    from datetime import date

    from apps.academics.services import create_academic_year

    other = create_academic_year(
        school, name="2027-2028", start_date=date(2027, 9, 1), end_date=date(2028, 6, 30), term_count=3
    )
    response = director.get(f"{REPORTS}subjects/", {"academic_year": other.pk, "term": term.pk})
    assert response.status_code == 400 and "term" in response.data["fields"]
    assert gradebook_of(setup["french"], term).status == "open"
