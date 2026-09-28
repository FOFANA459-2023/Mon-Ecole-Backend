import csv
import io

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def tabular_response(headers: list[str], rows: list[list], file_format: str, basename: str) -> HttpResponse:
    """Download rows as an Excel workbook (default) or a CSV file that opens cleanly in Excel."""
    stamp = timezone.localdate().isoformat()
    if file_format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(headers)
        writer.writerows(rows)
        response = HttpResponse("﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = f'attachment; filename="{basename}-{stamp}.csv"'
        return response

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = basename[:31]
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    for index, header in enumerate(headers, start=1):
        longest = max([len(str(header))] + [len(str(r[index - 1] or "")) for r in rows[:500]])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 45)

    buffer = io.BytesIO()
    workbook.save(buffer)
    response = HttpResponse(
        buffer.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{basename}-{stamp}.xlsx"'
    return response
