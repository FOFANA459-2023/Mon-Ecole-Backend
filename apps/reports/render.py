"""Report renderers: JSON for the screen, PDF to print, Excel and CSV to work on the figures."""

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from apps.core.pdf import LIGHT, MUTED, NAVY, format_date
from apps.finance.money import minor_places, to_money
from apps.finance.pdf import CELL, SUB, TITLE, document_header, format_money

from .base import Column, Report

PREVIEW_ROWS = 1000
NUMERIC = {"money", "number", "percent"}
GENERATED = {"fr": "Édité le {when} par {name}", "en": "Printed {when} by {name}"}
PAGE = {"fr": "Page {n}", "en": "Page {n}"}


def _local(value: datetime) -> datetime:
    return timezone.localtime(value) if timezone.is_aware(value) else value


# --- JSON ---------------------------------------------------------------------------------------------------


def _json_value(value, kind: str, currency: str):
    if value is None:
        return None
    if kind == "money":
        return f"{to_money(value, currency):.2f}"
    if kind == "percent":
        return str(value)
    if isinstance(value, datetime):
        return _local(value).isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _json_row(columns: list[Column], row: dict, currency: str, *, only_given: bool = False) -> dict:
    """`only_given`: a totals row keeps just the columns it totals, so the others stay blank (not "—")."""
    out = {
        column.key: _json_value(row.get(column.key), column.kind_for(row), currency)
        for column in columns
        if not only_given or column.key in row
    }
    if "kind" in row:
        out["kind"] = row["kind"]
    return out


def to_json(report: Report, currency: str) -> dict:
    """The report for the screen; long sections are cut to PREVIEW_ROWS rows (the exports have them all)."""
    return {
        "key": report.key,
        "title": report.title,
        "subtitle": report.subtitle,
        "currency": currency,
        "sections": [
            {
                "title": section.title,
                "note": section.note,
                "columns": [{"key": c.key, "label": c.label, "kind": c.kind} for c in section.columns],
                "rows": [_json_row(section.columns, row, currency) for row in section.rows[:PREVIEW_ROWS]],
                "row_count": len(section.rows),
                "truncated": len(section.rows) > PREVIEW_ROWS,
                "totals": (
                    _json_row(section.columns, section.totals, currency, only_given=True)
                    if section.totals
                    else None
                ),
            }
            for section in report.sections
        ],
    }


# --- CSV ----------------------------------------------------------------------------------------------------


def _plain(value, kind: str, currency: str) -> str:
    if value is None or value == "":
        return ""
    if kind == "money":
        return str(to_money(value, currency))
    if isinstance(value, datetime):
        return _local(value).strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def to_csv(report: Report, currency: str, language: str) -> str:
    """French Excel splits columns on ";" and English Excel on ","; amounts keep a "." decimal point."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";" if language == "fr" else ",")
    writer.writerow([report.title])
    writer.writerow([report.subtitle])
    for section in report.sections:
        writer.writerow([])
        writer.writerow([section.title])
        writer.writerow([column.label for column in section.columns])
        for row in [*section.rows, *([section.totals] if section.totals else [])]:
            writer.writerow([_plain(row.get(c.key), c.kind_for(row), currency) for c in section.columns])
    return "﻿" + buffer.getvalue()


# --- Excel --------------------------------------------------------------------------------------------------


def _sheet_title(title: str, taken: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", title).strip()[:28] or "Report"
    name, n = base, 2
    while name.lower() in taken:
        name, n = f"{base[:25]} {n}", n + 1
    taken.add(name.lower())
    return name


def to_xlsx(report: Report, currency: str) -> bytes:
    """One sheet per section, with real numbers and dates so the figures can be summed and sorted."""
    money_format = "#,##0" if minor_places(currency) == 0 else "#,##0.00"
    header_fill = PatternFill("solid", fgColor="1F3A5F")
    total_border = Border(top=Side(style="thin", color="1F3A5F"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    taken: set[str] = set()
    for section in report.sections:
        sheet = workbook.create_sheet(_sheet_title(section.title, taken))
        sheet.append([report.title])
        sheet["A1"].font = Font(bold=True, size=14, color="1F3A5F")
        sheet.append([report.subtitle])
        sheet.append([section.title + (f" — {section.note}" if section.note else "")])
        sheet["A3"].font = Font(bold=True)
        sheet.append([column.label for column in section.columns])
        header_row = sheet.max_row
        for cell in sheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
        sheet.freeze_panes = f"A{header_row + 1}"
        rows = [(row, False) for row in section.rows] + ([(section.totals, True)] if section.totals else [])
        for row, is_total in rows:
            values = []
            for column in section.columns:
                value, kind = row.get(column.key), column.kind_for(row)
                if value is not None and kind == "money":
                    value = float(to_money(value, currency))
                elif value is not None and kind == "percent":
                    value = float(value) / 100
                elif isinstance(value, datetime):
                    value = _local(value).replace(tzinfo=None)
                values.append(value)
            sheet.append(values)
            for column, cell in zip(section.columns, sheet[sheet.max_row], strict=True):
                kind = column.kind_for(row)
                if kind == "money":
                    cell.number_format = money_format
                elif kind == "percent":
                    cell.number_format = "0.0%"
                elif kind == "date":
                    cell.number_format = "dd/mm/yyyy"
                elif kind == "datetime":
                    cell.number_format = "dd/mm/yyyy hh:mm"
                if is_total:
                    cell.font = Font(bold=True)
                    cell.border = total_border
        for index, column in enumerate(section.columns, start=1):
            longest = max(
                [len(column.label)]
                + [len(_plain(r.get(column.key), column.kind_for(r), currency)) for r in section.rows[:300]]
            )
            sheet.column_dimensions[get_column_letter(index)].width = min(max(10, longest + 2), 45)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


# --- PDF ----------------------------------------------------------------------------------------------------

SMALL = ParagraphStyle("small", parent=CELL, fontSize=7.5, leading=9.5)
SMALL_NUM = ParagraphStyle("smallnum", parent=SMALL, alignment=2)
SMALL_BOLD = ParagraphStyle("smallbold", parent=SMALL, fontName="Helvetica-Bold")
SMALL_BOLD_NUM = ParagraphStyle("smallboldnum", parent=SMALL_NUM, fontName="Helvetica-Bold")
HEADING = ParagraphStyle("heading", parent=TITLE, fontSize=11, spaceBefore=4, spaceAfter=2)


def _display(value, kind: str, currency: str, school, language: str) -> str:
    if value is None or value == "":
        return "—" if kind in NUMERIC else ""
    if kind == "money":
        return format_money(to_money(value, currency), currency, language)
    if kind == "percent":
        return f"{value} %".replace(".", "," if language == "fr" else ".")
    if kind == "datetime" and isinstance(value, datetime):
        local = _local(value)
        return f"{format_date(local.date(), school)} {local:%H:%M}"
    if kind == "date" and isinstance(value, date):
        return format_date(value, school)
    return str(value)


CHAR_WIDTH = 1.45 * mm  # an average character at 7.5 pt
CELL_PADDING = 4 * mm


def _widths(section, texts: list[list[str]], width: float) -> list[float]:
    """Give each column the width of its longest content (capped), then shrink them to fit the page, or
    stretch them, but not so much that a small table's columns drift apart."""
    natural = []
    for index, column in enumerate(section.columns):
        longest = max([len(column.label)] + [len(row[index]) for row in texts[:200]])
        natural.append(min(longest, 40) * CHAR_WIDTH + CELL_PADDING)
    scale = min(width / sum(natural), 1.8)
    return [w * scale for w in natural]


def to_pdf(report: Report, school, language: str, printed_by: str) -> bytes:
    currency = school.currency
    pagesize = landscape(A4) if report.landscape else A4
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=report.title,
    )
    width = pagesize[0] - 24 * mm
    now = timezone.localtime()
    printed = GENERATED.get(language, GENERATED["fr"]).format(
        when=f"{format_date(now.date(), school)} {now:%H:%M}", name=printed_by
    )
    story = [
        document_header(school, report.title, "", "", escape(report.subtitle), width, 16),
        Spacer(1, 4 * mm),
    ]

    for section in report.sections:
        texts = [
            [_display(row.get(c.key), c.kind_for(row), currency, school, language) for c in section.columns]
            for row in section.rows
        ]
        header = [
            Paragraph(f"<b>{escape(c.label)}</b>", SMALL_NUM if c.kind in NUMERIC else SMALL)
            for c in section.columns
        ]
        body = [
            [
                Paragraph(escape(text), SMALL_NUM if column.kind_for(row) in NUMERIC else SMALL)
                for column, text in zip(section.columns, texts_row, strict=True)
            ]
            for row, texts_row in zip(section.rows, texts, strict=True)
        ]
        if section.totals:
            body.append(
                [
                    Paragraph(
                        escape(
                            _display(
                                section.totals.get(c.key),
                                c.kind_for(section.totals),
                                currency,
                                school,
                                language,
                            )
                        )
                        if c.key in section.totals
                        else "",
                        SMALL_BOLD_NUM if c.kind in NUMERIC else SMALL_BOLD,
                    )
                    for c in section.columns
                ]
            )
        table = Table([header, *body], colWidths=_widths(section, texts, width), repeatRows=1, hAlign="LEFT")
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), LIGHT),
            ("LINEBELOW", (0, 0), (-1, -1), 0.3, LIGHT),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        if section.totals:
            style.append(("LINEABOVE", (0, -1), (-1, -1), 0.8, NAVY))
        table.setStyle(TableStyle(style))
        heading = [Paragraph(escape(section.title), HEADING)]
        if section.note:
            heading.append(Paragraph(escape(section.note), SUB))
        if not section.rows:
            heading.append(Paragraph("—", SUB))
            story += [KeepTogether(heading), Spacer(1, 4 * mm)]
            continue
        # Keep a section's title with the start of its table.
        story += [KeepTogether([*heading, Spacer(1, 1 * mm), table]) if len(body) < 25 else heading[0]]
        if len(body) >= 25:
            story += [*heading[1:], Spacer(1, 1 * mm), table]
        story.append(Spacer(1, 5 * mm))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(12 * mm, 7 * mm, f"{school.name} — {report.title} — {printed}")
        canvas.drawRightString(
            pagesize[0] - 12 * mm, 7 * mm, PAGE.get(language, PAGE["fr"]).format(n=document.page)
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
