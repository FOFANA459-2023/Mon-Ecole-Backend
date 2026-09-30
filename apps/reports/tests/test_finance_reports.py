import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.accounts.models import Membership, Role
from apps.audit.models import AuditLog
from apps.enrollments import services as enrollment_services
from apps.finance import services as finance
from apps.people.services import link_guardian

URL = "/api/v1/reports/finance/{}/"
SEPTEMBER = {"date_from": "2026-09-01", "date_to": "2026-09-30"}


@pytest.fixture
def ledger(school, fees, make_class, make_student, cash_session):
    """September: Awa pays 1 250 000 in cash and 200 000 by mobile money (reversed); Binta pays nothing.
    100 000 spent on chalk, 50 000 of Awa's later overpayment refunded."""
    awa = make_student("Awa", "Diallo")
    binta = make_student("Binta", "Bah")
    klass = make_class()
    for student in (awa, binta):
        enrollment_services.enroll(student, klass, enrollment_date=date(2026, 9, 2))
    link_guardian(
        binta,
        guardian_data={"first_name": "Mariama", "last_name": "Bah", "phone": "620112233"},
        is_financial_contact=True,
    )
    finance.record_payment(awa, amount=Decimal("1250000"), method="cash", payment_date=date(2026, 9, 10))
    mistake = finance.record_payment(
        awa, amount=Decimal("200000"), method="mobile_money", payment_date=date(2026, 9, 11)
    )
    finance.reverse_payment(mistake, reason="Wrong student")
    finance.record_expense(
        school,
        amount=Decimal("100000"),
        category="supplies",
        method="cash",
        description="Chalk",
        expense_date=date(2026, 9, 12),
    )
    finance.record_payment(
        awa, amount=Decimal("2050000"), method="bank_transfer", payment_date=date(2026, 9, 20)
    )
    finance.record_refund(
        awa, amount=Decimal("50000"), method="cash", reason="Overpaid", refund_date=date(2026, 9, 21)
    )
    return {"awa": awa, "binta": binta, "class": klass}


@pytest.fixture
def accountant(school, make_member, client_for):
    return client_for(make_member(school, "accountant", language="en"), school)


def _section(data, index=0):
    return data["sections"][index]


@pytest.mark.django_db
class TestReports:
    def test_payments_report_lists_reversals_but_does_not_count_them(self, ledger, accountant):
        data = accountant.get(URL.format("payments"), SEPTEMBER).data
        assert data["title"] == "Payments report"
        payments = _section(data)
        assert [row["amount"] for row in payments["rows"]] == ["1250000.00", "200000.00", "2050000.00"]
        assert [row["status"] for row in payments["rows"]] == ["Recorded", "Reversed", "Recorded"]
        assert payments["rows"][0]["class"] == "7ème A"
        assert payments["totals"]["amount"] == "3300000.00"
        by_method = _section(data, 1)
        assert {row["method"]: row["amount"] for row in by_method["rows"]} == {
            "Bank transfer / deposit": "2050000.00",
            "Cash": "1250000.00",
        }

    def test_filters_and_period(self, ledger, accountant):
        cash_only = accountant.get(URL.format("payments"), {**SEPTEMBER, "method": "cash"}).data
        assert len(_section(cash_only)["rows"]) == 1
        early = accountant.get(
            URL.format("payments"), {"date_from": "2026-09-01", "date_to": "2026-09-10"}
        ).data
        assert _section(early)["totals"]["amount"] == "1250000.00"

    def test_outstanding_balances_with_the_guardian_to_call(self, ledger, accountant):
        data = accountant.get(URL.format("outstanding")).data
        rows = _section(data)["rows"]
        # Awa paid everything (3 250 000); Binta owes it all.
        assert [(row["student"], row["balance"]) for row in rows] == [("Binta Bah", "3250000.00")]
        assert (rows[0]["guardian"], rows[0]["phone"]) == ("Mariama Bah", "620112233")
        assert _section(data)["totals"]["balance"] == "3250000.00"

    def test_cash_report_day_by_day(self, ledger, accountant, cash_session):
        data = accountant.get(URL.format("cash"), SEPTEMBER).data
        days, sessions, by_source = data["sections"]
        assert [row["date"] for row in days["rows"]] == ["2026-09-01"]
        # In: 1 250 000 payment. Out: 100 000 chalk + 50 000 refund.
        assert (sessions["totals"]["money_in"], sessions["totals"]["money_out"]) == (
            "1250000.00",
            "150000.00",
        )
        assert sessions["rows"][0]["expected"] == "1100000.00"
        assert {(row["source"], row["amount"]) for row in by_source["rows"]} == {
            ("Payments", "1250000.00"),
            ("Expenses", "100000.00"),
            ("Refunds", "50000.00"),
        }

    def test_expenses_report(self, ledger, accountant, school):
        cancelled = finance.record_expense(
            school,
            amount=Decimal("5000"),
            category="other",
            method="mobile_money",
            description="Typo",
            expense_date=date(2026, 9, 13),
        )
        finance.cancel_expense(cancelled, reason="Typo")
        data = accountant.get(URL.format("expenses"), SEPTEMBER).data
        assert len(_section(data)["rows"]) == 2
        assert _section(data)["totals"]["amount"] == "100000.00"
        assert _section(data, 1)["rows"] == [{"category": "Supplies", "count": 1, "amount": "100000.00"}]

    def test_financial_summary(self, ledger, accountant, year):
        data = accountant.get(URL.format("summary"), SEPTEMBER).data
        overview = _section(data)
        assert [row["amount"] for row in overview["rows"]] == ["3300000.00", "-50000.00", "-100000.00"]
        assert overview["totals"]["amount"] == "3150000.00"
        by_fee = {row["fee"]: row["amount"] for row in _section(data, 2)["rows"]}
        assert by_fee == {
            "Scolarité": "3000000.00",
            "Inscription": "250000.00",
            "Advances kept as credit": "50000.00",
        }
        months = _section(data, 4)["rows"]
        assert months == [
            {
                "month": "Sep 2026",
                "received": "3300000.00",
                "refunds": "50000.00",
                "spent": "100000.00",
                "net": "3150000.00",
            }
        ]
        position = {row["item"]: (row["value"], row["kind"]) for row in _section(data, 5)["rows"]}
        assert position["Expected revenue (invoices issued)"] == ("6500000.00", "money")
        assert position["Collection rate"] == ("50.0", "percent")
        assert position["Students with a balance due"] == (1, "number")

    def test_french_labels_follow_the_user(self, ledger, school, make_member, client_for):
        client = client_for(make_member(school, "director", language="fr"), school)
        data = client.get(URL.format("payments"), SEPTEMBER).data
        assert data["title"] == "Rapport des paiements"
        assert data["subtitle"] == "Du 01/09/2026 au 30/09/2026"


@pytest.mark.django_db
class TestExports:
    def test_pdf(self, ledger, accountant):
        response = accountant.get(URL.format("outstanding"), {"export": "pdf"})
        assert response.status_code == 200
        assert response["Content-Type"] == "application/pdf"
        assert response.content.startswith(b"%PDF")

    def test_excel_keeps_real_numbers(self, ledger, accountant):
        response = accountant.get(URL.format("payments"), {**SEPTEMBER, "export": "xlsx"})
        assert response.status_code == 200
        assert "attachment" in response["Content-Disposition"]
        workbook = load_workbook(io.BytesIO(response.content))
        assert workbook.sheetnames == ["Payments", "By payment method"]
        sheet = workbook["Payments"]
        header = [cell.value for cell in sheet[4]]
        amount_column = header.index("Amount") + 1
        assert sheet.cell(row=5, column=amount_column).value == 1250000
        assert sheet.cell(row=5, column=amount_column).number_format == "#,##0"

    def test_csv_uses_the_separator_excel_expects(self, ledger, school, accountant, make_member, client_for):
        english = accountant.get(URL.format("expenses"), {**SEPTEMBER, "export": "csv"}).content.decode(
            "utf-8-sig"
        )
        assert "Chalk" in english
        assert next(csv.reader(io.StringIO(english.splitlines()[4])))[:2] == ["No.", "Date"]
        french = client_for(make_member(school, "director", language="fr"), school)
        text = french.get(URL.format("expenses"), {**SEPTEMBER, "export": "csv"}).content.decode("utf-8-sig")
        assert "N°;Date;Catégorie" in text

    def test_exports_are_audited(self, ledger, accountant):
        accountant.get(URL.format("summary"), {**SEPTEMBER, "export": "pdf"})
        log = AuditLog.objects.get(action="export")
        assert (log.entity_type, log.entity_id, log.new_values["format"]) == ("report", "summary", "pdf")

    def test_exporting_needs_the_export_right(self, school, ledger, make_member, client_for):
        viewer = make_member(school, "teacher", email="viewer@test.local")
        role = Role.objects.create(school=school, name="Finance viewer", permissions=["finance.view"])
        Membership.objects.get(user=viewer).roles.set([role])
        client = client_for(viewer, school)
        assert client.get(URL.format("payments")).status_code == 200
        assert client.get(URL.format("payments"), {"export": "csv"}).status_code == 403


@pytest.mark.django_db
class TestAccess:
    def test_who_may_read_reports(self, school, make_member, client_for):
        for role in ("teacher", "admin_staff"):
            assert (
                client_for(make_member(school, role), school).get(URL.format("payments")).status_code == 403
            )

    def test_unknown_report_and_bad_periods(self, school, accountant):
        assert accountant.get(URL.format("nonsense")).status_code == 404
        backwards = accountant.get(
            URL.format("payments"), {"date_from": "2026-09-30", "date_to": "2026-09-01"}
        )
        assert backwards.status_code == 400
        too_long = accountant.get(
            URL.format("payments"), {"date_from": "2020-01-01", "date_to": "2026-09-01"}
        )
        assert too_long.status_code == 400

    def test_only_the_school_own_figures(self, ledger, other_school, make_member, client_for):
        outsider = client_for(make_member(other_school, "accountant"), other_school)
        assert _section(outsider.get(URL.format("payments"), SEPTEMBER).data)["rows"] == []
        assert _section(outsider.get(URL.format("outstanding")).data)["rows"] == []
        # A class of another school is refused, not silently ignored.
        response = outsider.get(URL.format("outstanding"), {"class_group": ledger["class"].pk})
        assert response.status_code == 400
